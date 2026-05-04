"""
单元测试：ExcelGenerator

测试 PO 合同 Excel 文件的生成逻辑，包括：
- 文件成功创建
- 工作表内容（标题、供应商信息、行项目合计）
- 文件命名规则
"""
import os
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

# 让测试可以导入项目根模块
sys.path.insert(0, str(Path(__file__).parent.parent))

from excel.generator import ExcelGenerator


@pytest.fixture
def tmp_output(tmp_path: Path) -> Path:
    return tmp_path / "contracts"


@pytest.fixture
def sample_po_data() -> dict:
    return {
        "_record_id": "rec001",
        "po_number": "PO-2024-001",
        "contract_date": "2024-01-15",
        "supplier_name": "测试供应商有限公司",
        "supplier_address": "上海市浦东新区XX路100号",
        "supplier_contact": "张三",
        "supplier_tel": "021-12345678",
        "buyer_name": "测试采购公司",
        "buyer_dept": "采购部",
        "buyer_contact": "李四",
        "currency": "CNY",
        "payment_terms": "月结30天",
        "delivery_date": "2024-02-01",
        "delivery_address": "北京市朝阳区YY街道200号",
        "items": [
            {"description": "产品A", "unit": "个", "qty": 10, "unit_price": 100.0},
            {"description": "产品B", "unit": "套", "qty": 5, "unit_price": 250.0},
        ],
        "notes": "请确保按时交货",
    }


class TestExcelGeneratorFileCreation:
    """测试文件生成"""

    def test_generates_xlsx_file(self, tmp_output: Path, sample_po_data: dict):
        gen = ExcelGenerator(output_dir=str(tmp_output))
        filepath = gen.generate(sample_po_data)
        assert Path(filepath).exists(), "生成的 Excel 文件应存在"
        assert filepath.endswith(".xlsx"), "文件扩展名应为 .xlsx"

    def test_creates_output_directory_if_missing(self, tmp_output: Path, sample_po_data: dict):
        nested = tmp_output / "nested" / "dir"
        gen = ExcelGenerator(output_dir=str(nested))
        gen.generate(sample_po_data)
        assert nested.exists(), "应自动创建输出目录"

    def test_filename_contains_po_number(self, tmp_output: Path, sample_po_data: dict):
        gen = ExcelGenerator(output_dir=str(tmp_output))
        filepath = gen.generate(sample_po_data)
        filename = os.path.basename(filepath)
        assert "PO-2024-001" in filename, "文件名应包含 PO 编号"

    def test_filename_prefix_is_po(self, tmp_output: Path, sample_po_data: dict):
        gen = ExcelGenerator(output_dir=str(tmp_output))
        filepath = gen.generate(sample_po_data)
        filename = os.path.basename(filepath)
        assert filename.startswith("PO_"), "文件名应以 PO_ 开头"


class TestExcelGeneratorContent:
    """测试 Excel 工作表内容"""

    @pytest.fixture(autouse=True)
    def setup(self, tmp_output: Path, sample_po_data: dict):
        gen = ExcelGenerator(output_dir=str(tmp_output))
        filepath = gen.generate(sample_po_data)
        self.wb = load_workbook(filepath)
        self.ws = self.wb.active
        self.po_data = sample_po_data

    def test_sheet_name(self):
        assert self.ws.title == "采购订单"

    def _all_values(self) -> list:
        return [
            cell.value
            for row in self.ws.iter_rows()
            for cell in row
            if cell.value is not None
        ]

    def test_company_name_in_header(self):
        from config import CompanyConfig
        values = self._all_values()
        assert CompanyConfig.NAME in values, "表头应包含公司名称"

    def test_po_title_present(self):
        values = self._all_values()
        assert any("采" in str(v) and "购" in str(v) and "订" in str(v) and "单" in str(v)
                   for v in values), "工作表应包含'采购订单'字样"

    def test_supplier_name_present(self):
        values = self._all_values()
        assert "测试供应商有限公司" in values, "供应商名称应出现在工作表中"

    def test_buyer_name_present(self):
        values = self._all_values()
        assert "测试采购公司" in values, "采购方名称应出现在工作表中"

    def test_po_number_present(self):
        values = self._all_values()
        assert "PO-2024-001" in values, "PO 编号应出现在工作表中"

    def test_item_total_calculated_correctly(self):
        """合计金额 = 10*100 + 5*250 = 2250"""
        expected_total = 10 * 100.0 + 5 * 250.0
        values = self._all_values()
        assert expected_total in values, f"合计金额应为 {expected_total}"

    def test_item_description_present(self):
        values = self._all_values()
        assert "产品A" in values, "行项目描述应出现在工作表中"

    def test_payment_terms_present(self):
        values = self._all_values()
        assert "月结30天" in values, "付款条款应出现在工作表中"


class TestExcelGeneratorEdgeCases:
    """测试边界情况"""

    def test_empty_items_list(self, tmp_output: Path):
        """items 为空时也应成功生成文件"""
        po_data = {
            "po_number": "PO-EMPTY",
            "contract_date": "2024-01-01",
            "supplier_name": "供应商",
            "buyer_name": "采购方",
            "currency": "CNY",
            "items": [],
        }
        gen = ExcelGenerator(output_dir=str(tmp_output))
        filepath = gen.generate(po_data)
        assert Path(filepath).exists()

    def test_missing_optional_fields(self, tmp_output: Path):
        """可选字段缺失时应正常运行"""
        po_data = {"po_number": "PO-MIN", "items": []}
        gen = ExcelGenerator(output_dir=str(tmp_output))
        filepath = gen.generate(po_data)
        assert Path(filepath).exists()

    def test_po_number_with_slash_sanitized(self, tmp_output: Path):
        """PO 编号含斜杠时，文件名中应替换为连字符"""
        po_data = {"po_number": "PO/2024/001", "items": []}
        gen = ExcelGenerator(output_dir=str(tmp_output))
        filepath = gen.generate(po_data)
        assert "/" not in os.path.basename(filepath), "文件名中不应包含斜杠"

    def test_items_with_zero_price(self, tmp_output: Path):
        """单价为 0 时合计应为 0"""
        po_data = {
            "po_number": "PO-ZERO",
            "items": [{"description": "测试品", "qty": 5, "unit_price": 0}],
        }
        gen = ExcelGenerator(output_dir=str(tmp_output))
        filepath = gen.generate(po_data)
        wb = load_workbook(filepath)
        ws = wb.active
        values = [
            cell.value
            for row in ws.iter_rows()
            for cell in row
            if cell.value is not None
        ]
        assert 0.0 in values or 0 in values
